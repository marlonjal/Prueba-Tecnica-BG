"""Unit tests for the styled Excel Pytest result export."""

from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

from core.result_exporter import _status_and_detail, _write_xlsx


def test_status_export_distinguishes_passed_and_expected_failure() -> None:
    """Verify the exported labels for a pass and a known functional defect."""
    passed = SimpleNamespace(
        when="call", passed=True, failed=False, skipped=False
    )
    xfailed = SimpleNamespace(
        when="call",
        passed=False,
        failed=False,
        skipped=True,
        wasxfail="known ParaBank defect",
    )

    assert _status_and_detail(passed) == ("APROBADA", "")
    assert _status_and_detail(xfailed) == ("XFAIL", "known ParaBank defect")


def test_xlsx_export_contains_summary_and_filterable_detail(tmp_path: Path) -> None:
    """Verify workbook sheets, accents, formulas, and the result table."""
    output = tmp_path / "test-results.xlsx"
    _write_xlsx(
        output,
        [
            {
                "ID": 1,
                "TIPO": "BDD",
                "ARCHIVO": "test_login_steps.py",
                "PRUEBA": "test_inicio_de_sesión_exitoso",
                "ESTADO": "APROBADA",
                "DURACION_SEGUNDOS": 1.25,
                "DETALLE": "",
            }
        ],
    )

    workbook = load_workbook(output, data_only=False)

    assert workbook.sheetnames == ["Resumen", "Detalle"]
    assert workbook["Detalle"]["D6"].value == "test_inicio_de_sesión_exitoso"
    assert workbook["Detalle"].tables["ResultadosPruebas"].ref == "A5:G6"
    assert workbook["Resumen"]["B5"].value.startswith("=COUNTIF(")

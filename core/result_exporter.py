"""Export every selected Pytest result to a styled Excel workbook."""

from __future__ import annotations

from collections import Counter
from datetime import datetime
from pathlib import Path, PurePath

import pytest
from openpyxl import Workbook
from openpyxl.chart import DoughnutChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

_reports: dict[str, pytest.TestReport] = {}
_durations: Counter[str] = Counter()
_node_order: list[str] = []

_FIELDNAMES = (
    "ID",
    "TIPO",
    "ARCHIVO",
    "PRUEBA",
    "ESTADO",
    "DURACION_SEGUNDOS",
    "DETALLE",
)

_NAVY = "17365D"
_BLUE = "1F4E78"
_LIGHT_BLUE = "D9EAF7"
_WHITE = "FFFFFF"
_TEXT = "1F2937"
_BORDER = Side(style="thin", color="D9E2F3")
_STATUS_COLORS = {
    "APROBADA": "C6EFCE",
    "FALLIDA": "FFC7CE",
    "ERROR": "F4CCCC",
    "OMITIDA": "FFF2CC",
    "XFAIL": "FCE4D6",
    "XPASS": "D9EAD3",
    "DESCONOCIDO": "E7E6E6",
}


def pytest_collection_finish(session: pytest.Session) -> None:
    """Reset exporter state and preserve the selected test execution order."""
    _reports.clear()
    _durations.clear()
    _node_order.clear()
    _node_order.extend(item.nodeid for item in session.items)


def pytest_runtest_logreport(report: pytest.TestReport) -> None:
    """Accumulate phase duration and retain the final meaningful test report."""
    _durations[report.nodeid] += report.duration

    if report.when == "call":
        _reports[report.nodeid] = report
    elif report.when in {"setup", "teardown"} and (report.failed or report.skipped):
        _reports[report.nodeid] = report


def pytest_sessionfinish(session: pytest.Session) -> None:
    """Write the selected execution results after Pytest finishes all tests."""
    if session.config.option.collectonly:
        return

    output_path = Path(str(session.config.rootpath)) / "results" / "reports"
    rows = [
        _result_row(position, nodeid, _reports[nodeid])
        for position, nodeid in enumerate(_node_order, start=1)
        if nodeid in _reports
    ]
    _write_xlsx(output_path / "test-results.xlsx", rows)
    (output_path / "test-results.csv").unlink(missing_ok=True)


def _result_row(
    position: int, nodeid: str, report: pytest.TestReport
) -> dict[str, object]:
    """Transform a Pytest report into the stable Spanish export schema."""
    file_path, *test_parts = nodeid.split("::")
    status, detail = _status_and_detail(report)
    normalized_path = file_path.replace("\\", "/")
    return {
        "ID": position,
        "TIPO": "BDD" if "/steps/" in f"/{normalized_path}" else "UNITARIA",
        "ARCHIVO": PurePath(normalized_path).name,
        "PRUEBA": "::".join(test_parts),
        "ESTADO": status,
        "DURACION_SEGUNDOS": round(_durations[nodeid], 3),
        "DETALLE": detail,
    }


def _status_and_detail(report: pytest.TestReport) -> tuple[str, str]:
    """Normalize native Pytest outcomes and their most useful explanation."""
    if hasattr(report, "wasxfail"):
        status = "XFAIL" if report.skipped else "XPASS"
        return status, str(report.wasxfail)
    if report.when in {"setup", "teardown"} and report.failed:
        return "ERROR", _failure_detail(report)
    if report.passed:
        return "APROBADA", ""
    if report.skipped:
        return "OMITIDA", _skipped_detail(report)
    if report.failed:
        return "FALLIDA", _failure_detail(report)
    return "DESCONOCIDO", ""


def _failure_detail(report: pytest.TestReport) -> str:
    """Return the final diagnostic line from a failed Pytest report."""
    lines = report.longreprtext.strip().splitlines()
    return lines[-1] if lines else ""


def _skipped_detail(report: pytest.TestReport) -> str:
    """Return the reason attached to a skipped or unavailable test."""
    if isinstance(report.longrepr, tuple) and len(report.longrepr) >= 3:
        return str(report.longrepr[2])
    return str(report.longrepr)


def _write_xlsx(path: Path, rows: list[dict[str, object]]) -> None:
    """Create a styled workbook with an executive summary and test details."""
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumen"
    details = workbook.create_sheet("Detalle")

    _build_detail_sheet(details, rows)
    _build_summary_sheet(summary, len(rows), details.max_row)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(path)


def _build_detail_sheet(sheet, rows: list[dict[str, object]]) -> None:
    """Populate the filterable detail table and color every execution status."""
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:G2")
    title = sheet["A1"]
    title.value = "DETALLE DE RESULTADOS DE AUTOMATIZACIÓN"
    title.fill = PatternFill("solid", fgColor=_NAVY)
    title.font = Font(color=_WHITE, bold=True, size=16)
    title.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 25
    sheet.row_dimensions[2].height = 10
    sheet["A3"] = f"Generado: {datetime.now().astimezone():%Y-%m-%d %H:%M:%S %Z}"
    sheet["A3"].font = Font(color="667085", italic=True, size=10)

    header_row = 5
    for column, field in enumerate(_FIELDNAMES, start=1):
        cell = sheet.cell(row=header_row, column=column, value=field)
        cell.fill = PatternFill("solid", fgColor=_BLUE)
        cell.font = Font(color=_WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_number, result in enumerate(rows, start=header_row + 1):
        for column, field in enumerate(_FIELDNAMES, start=1):
            value = _safe_excel_value(result[field])
            cell = sheet.cell(row=row_number, column=column, value=value)
            cell.font = Font(color=_TEXT, size=10)
            cell.alignment = Alignment(
                horizontal="right" if field in {"ID", "DURACION_SEGUNDOS"} else "left",
                vertical="top",
                wrap_text=field in {"PRUEBA", "DETALLE"},
            )
            cell.border = Border(bottom=_BORDER)

        status_cell = sheet.cell(row=row_number, column=5)
        status_cell.fill = PatternFill(
            "solid", fgColor=_STATUS_COLORS.get(str(status_cell.value), "E7E6E6")
        )
        status_cell.font = Font(color=_TEXT, bold=True)
        status_cell.alignment = Alignment(horizontal="center")

    last_row = max(header_row + 1, header_row + len(rows))
    if not rows:
        sheet.cell(row=last_row, column=1, value="Sin pruebas ejecutadas")
    table = Table(displayName="ResultadosPruebas", ref=f"A{header_row}:G{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    sheet.freeze_panes = "A6"
    sheet.auto_filter.ref = f"A{header_row}:G{last_row}"
    sheet.column_dimensions["A"].width = 8
    sheet.column_dimensions["B"].width = 13
    sheet.column_dimensions["C"].width = 32
    sheet.column_dimensions["D"].width = 62
    sheet.column_dimensions["E"].width = 16
    sheet.column_dimensions["F"].width = 20
    sheet.column_dimensions["G"].width = 75
    sheet.column_dimensions["F"].number_format = "0.000"
    sheet.auto_filter.ref = table.ref


def _build_summary_sheet(sheet, total_rows: int, detail_last_row: int) -> None:
    """Build formula-driven totals and a status-distribution chart."""
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:H2")
    title = sheet["A1"]
    title.value = "REPORTE DE EJECUCIÓN AUTOMATIZADA"
    title.fill = PatternFill("solid", fgColor=_NAVY)
    title.font = Font(color=_WHITE, bold=True, size=18)
    title.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 28

    sheet["A4"] = "ESTADO"
    sheet["B4"] = "CANTIDAD"
    for cell in sheet[4][0:2]:
        cell.fill = PatternFill("solid", fgColor=_BLUE)
        cell.font = Font(color=_WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center")

    statuses = ["APROBADA", "FALLIDA", "ERROR", "OMITIDA", "XFAIL", "XPASS"]
    detail_range = f"'Detalle'!$E$6:$E${max(6, detail_last_row)}"
    for row_number, status in enumerate(statuses, start=5):
        sheet.cell(row=row_number, column=1, value=status)
        sheet.cell(
            row=row_number,
            column=2,
            value=f'=COUNTIF({detail_range},A{row_number})',
        )
        for cell in sheet[row_number][0:2]:
            cell.border = Border(bottom=_BORDER)
        sheet.cell(row=row_number, column=1).fill = PatternFill(
            "solid", fgColor=_STATUS_COLORS[status]
        )

    sheet["A12"] = "TOTAL EJECUTADAS"
    sheet["B12"] = f"=SUM(B5:B10)"
    sheet["A13"] = "DURACIÓN TOTAL (S)"
    sheet["B13"] = f"=SUM('Detalle'!$F$6:$F${max(6, detail_last_row)})"
    sheet["A14"] = "FILAS EXPORTADAS"
    sheet["B14"] = total_rows
    for row_number in range(12, 15):
        sheet.cell(row=row_number, column=1).fill = PatternFill(
            "solid", fgColor=_LIGHT_BLUE
        )
        sheet.cell(row=row_number, column=1).font = Font(bold=True, color=_TEXT)
        sheet.cell(row=row_number, column=2).font = Font(bold=True, color=_TEXT)

    chart = DoughnutChart()
    chart.title = "Distribución de resultados"
    chart.style = 10
    chart.height = 8
    chart.width = 13
    data = Reference(sheet, min_col=2, min_row=4, max_row=10)
    labels = Reference(sheet, min_col=1, min_row=5, max_row=10)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(labels)
    chart.holeSize = 55
    chart.legend.position = "r"
    sheet.add_chart(chart, "D4")

    sheet.column_dimensions["A"].width = 25
    sheet.column_dimensions["B"].width = 16
    for column in "CDEFGH":
        sheet.column_dimensions[column].width = 14
    sheet.freeze_panes = "A4"
    sheet.conditional_formatting.add(
        "B5:B10",
        FormulaRule(formula=["B5>0"], fill=PatternFill("solid", fgColor="E2F0D9")),
    )


def _safe_excel_value(value: object) -> object:
    """Prevent diagnostic text from being interpreted as an Excel formula."""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value
